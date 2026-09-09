"""현장 CCTV 영상에서 파인튜닝용 결함 프레임을 자동으로 뽑는다.

원리 — 조사 문서(`현리산유 하자목록.xlsx`)는 결함 위치를 **미터 단위**로 적어두고,
영상 좌상단에는 주행거리 자막이 박혀 있다. 프레임별 거리를 OCR로 읽으면
"42.5m 이음부이탈" 같은 문서 기록을 실제 프레임에 자동으로 붙일 수 있다.
영상을 눈으로 훑으며 결함을 찾는 작업이 사라지고, 박스만 그리면 된다.

거리는 단조증가하므로 전 프레임을 OCR하지 않는다. 성긴 간격으로 시간↔거리
곡선을 만들고(coarse), 목표 거리 부근만 촘촘히 다시 읽어(refine) 프레임을 고른다.

검증용 영상(`검증용 영상data`)과 같은 관로는 **자동으로 제외**한다. 그쪽은
최종 테스트셋이라 학습에 섞이면 안 된다.

출력: out_dir/{frames,manifest.csv}. manifest.csv를 검수한 뒤 박스를 그린다.
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from bisect import bisect_left
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import cv2
import numpy as np
import openpyxl

APP_DIR = Path(__file__).resolve().parents[2] / "apps" / "AI_CCTV"
if str(APP_DIR) not in sys.path:
    sys.path.insert(0, str(APP_DIR))

# 문서의 한글 결함명 -> 마스터 결함 코드(docs/메타데이터(총괄).xlsx).
# 앞쪽 항목부터 순서대로 검사하므로 긴 이름(더 구체적인 것)을 위에 둔다.
NAME_TO_CODE: List[Tuple[str, str]] = [
    ("이음부단차", "JD"), ("이음부 단차", "JD"),
    ("이음부이탈", "JS"), ("이음부 이탈", "JS"),
    ("이음부손상", "JF"), ("이음부 손상", "JF"),
    ("이음부파손", "JF"), ("이음부 파손", "JF"),
    ("이음부변형", "JF"), ("이음부 변형", "JF"),
    ("관접합부", "LS"),
    ("관주름", "DF"), ("관눌림", "DF"), ("관변형", "DF"), ("변형", "DF"),
    ("관침하", "SG"), ("침하", "SG"),
    ("관파손", "BK"), ("파손", "BK"),
    ("균열원주", "CC"), ("균열 원주", "CC"),
    ("균열길이", "CL"), ("균열 길이", "CL"),
    ("좌굴", "BC"),
    ("역구배", "NS"), ("역경사", "NS"),
    ("침입수", "IF"),
    ("토사", "DS"), ("퇴적", "DS"),
]
# 사람이 한 번 봐야 하는 애매한 표현
AMBIGUOUS = ("의심", "?", "검토", "확인")


@dataclass
class Defect:
    pipe: str
    pos_m: float
    raw: str
    code: str
    grade: str
    needs_review: bool


@dataclass
class VideoJob:
    pipe: str
    path: Path
    defects: List[Defect] = field(default_factory=list)


def classify(raw: str) -> Tuple[str, str, bool]:
    """문서의 하자유형 문자열 -> (코드, 등급, 검수필요)."""
    s = str(raw).strip()
    grade = ""
    m = re.search(r"[-\s]\s*(소|중|대)\s*$", s)
    if m:
        grade = m.group(1)
    code = ""
    flat = s.replace(" ", "")
    for name, c in NAME_TO_CODE:
        if name.replace(" ", "") in flat:
            code = c
            break
    review = (not code) or any(a in s for a in AMBIGUOUS) or "/" in s
    return code, grade, review


def load_defects(xlsx: Path) -> List[Defect]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out: List[Defect] = []
    for name in wb.sheetnames:
        if name in ("연장집계", "하자집계"):
            continue
        for r in wb[name].iter_rows(min_row=4, values_only=True):
            pipe = str(r[1]).strip() if r[1] else ""
            pos = r[18] if len(r) > 18 else None
            kind = r[19] if len(r) > 19 else None
            if not pipe or not isinstance(kind, str):
                continue
            k = kind.strip()
            if k in ("이상없음", "추가관로", ""):
                continue
            # 위치 칸이 비었거나 0인데 설명에 거리가 박혀 있는 경우
            # (예: 하자유형 "26m 침하") 설명에서 거리를 꺼낸다.
            if not isinstance(pos, (int, float)) or float(pos) == 0.0:
                m2 = re.search(r"(\d+(?:\.\d+)?)\s*m(?![/a-zA-Z])", k)
                if not m2:
                    continue
                pos = float(m2.group(1))
            code, grade, review = classify(k)
            out.append(Defect(pipe, float(pos), k, code, grade, review))
    return out


def norm_stem(name: str) -> str:
    s = Path(name).stem.lower()
    s = re.sub(r"^\d+\.\s*", "", s)
    s = re.sub(r"^masked_", "", s)
    s = re.sub(r"^r\s+", "", s)
    return s.strip()


def index_videos(root: Path) -> Dict[str, Path]:
    idx: Dict[str, Path] = {}
    for p in root.rglob("*.mp4"):
        if p.name.lower().startswith("masked_") or "masked" in p.parent.name.lower():
            continue
        idx.setdefault(norm_stem(p.name), p)
    return idx


def read_distance(ocr_fn, frame: np.ndarray, tmp: Path) -> Optional[float]:
    cv2.imencode(".jpg", frame, [int(cv2.IMWRITE_JPEG_QUALITY), 90])[1].tofile(str(tmp))
    txt = ocr_fn(tmp)
    if not txt:
        return None
    m = re.search(r"(\d+(?:\.\d+)?)", str(txt).replace(" ", ""))
    return float(m.group(1)) if m else None


def longest_nondecreasing(pts: List[Tuple[float, float]]) -> List[Tuple[float, float]]:
    """OCR 오독을 걸러낸다. 앞에서부터 탐욕적으로 버리면 첫 값이 잘못 읽혔을 때
    (예: 3.3을 33.0으로) 뒤가 통째로 날아가므로, 가장 긴 비감소 부분수열을 택한다."""
    n = len(pts)
    if n < 2:
        return pts
    best = [1] * n
    prev = [-1] * n
    for i in range(n):
        for j in range(i):
            if pts[j][1] <= pts[i][1] + 0.3 and best[j] + 1 > best[i]:
                best[i], prev[i] = best[j] + 1, j
    i = max(range(n), key=lambda k: best[k])
    out = []
    while i >= 0:
        out.append(pts[i])
        i = prev[i]
    return out[::-1]


def scan_curve(cap, ocr_fn, tmp: Path, fps: float, n_frames: int,
               step_s: float) -> List[Tuple[float, float]]:
    """성긴 간격으로 (시각, 거리) 곡선을 만든다."""
    raw: List[Tuple[float, float]] = []
    t = 0.0
    dur = n_frames / fps if fps else 0
    while t < dur:
        cap.set(cv2.CAP_PROP_POS_FRAMES, int(t * fps))
        ok, fr = cap.read()
        if not ok:
            break
        d = read_distance(ocr_fn, fr, tmp)
        if d is not None:
            raw.append((t, d))
        t += step_s
    return longest_nondecreasing(raw)


def time_for_distance(pts: List[Tuple[float, float]], target: float) -> Optional[float]:
    if len(pts) < 2:
        return None
    ds = [d for _, d in pts]
    if target < ds[0] - 3 or target > ds[-1] + 3:
        return None
    i = bisect_left(ds, target)
    i = max(1, min(i, len(pts) - 1))
    (t0, d0), (t1, d1) = pts[i - 1], pts[i]
    if d1 == d0:
        return t0
    return t0 + (t1 - t0) * (target - d0) / (d1 - d0)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=r"C:\Users\SAMSUNG\Desktop\CCTV\dataset\videos"
                                     r"\1. 1차\현리산유 하자목록(1차분-최종).xlsx")
    ap.add_argument("--videos", default=r"C:\Users\SAMSUNG\Desktop\CCTV\dataset\videos")
    ap.add_argument("--exclude", default=r"C:\Users\SAMSUNG\Desktop\CCTV\검증용 영상data",
                    help="테스트셋. 여기 있는 관로는 학습에서 제외한다")
    ap.add_argument("--out", default=r"E:\AI_CCTV_DATASET\fieldset_v1")
    ap.add_argument("--coarse-step", type=float, default=8.0, help="곡선 스캔 간격(초)")
    ap.add_argument("--tol-m", type=float, default=0.6, help="결함 위치 허용 오차(m)")
    ap.add_argument("--per-defect", type=int, default=5, help="결함당 뽑을 프레임 수")
    ap.add_argument("--normals-per-video", type=int, default=6)
    ap.add_argument("--limit-videos", type=int, default=0)
    args = ap.parse_args()

    from backend.ocr import ocr_distance_from_frame  # paddleocr 로딩이 느려 여기서 import

    out = Path(args.out)
    frames_dir = out / "frames"
    frames_dir.mkdir(parents=True, exist_ok=True)
    tmp = out / "_tmp.jpg"

    defects = load_defects(Path(args.xlsx))
    vids = index_videos(Path(args.videos))
    excluded = {norm_stem(p.name) for p in Path(args.exclude).glob("*.mp4")}

    jobs: Dict[str, VideoJob] = {}
    skipped_val, no_video = [], []
    for d in defects:
        key = norm_stem(d.pipe)
        if key in excluded:
            skipped_val.append(d)
            continue
        p = vids.get(key)
        if not p:
            no_video.append(d)
            continue
        jobs.setdefault(key, VideoJob(d.pipe, p)).defects.append(d)

    print(f"문서 결함 {len(defects)}건 | 검증용이라 제외 {len(skipped_val)} | "
          f"영상없음 {len(no_video)} | 처리대상 영상 {len(jobs)}개")

    job_list = sorted(jobs.values(), key=lambda j: j.pipe)
    if args.limit_videos:
        job_list = job_list[:args.limit_videos]

    rows = []
    for n, job in enumerate(job_list, 1):
        cap = cv2.VideoCapture(str(job.path))
        fps = cap.get(cv2.CAP_PROP_FPS) or 30.0
        nf = cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0
        pts = scan_curve(cap, ocr_distance_from_frame, tmp, fps, nf, args.coarse_step)
        span = f"{pts[0][1]:.1f}~{pts[-1][1]:.1f}m" if pts else "읽기실패"
        print(f"[{n}/{len(job_list)}] {job.pipe:<14} {job.path.name[:28]:<30} "
              f"곡선 {len(pts):>3}점 {span:<14} 결함 {len(job.defects)}건")
        if not pts:
            cap.release()
            continue

        got = []
        for d in job.defects:
            t0 = time_for_distance(pts, d.pos_m)
            if t0 is None:
                print(f"      - {d.pos_m:.1f}m {d.raw[:20]} : 거리범위 밖, 건너뜀")
                continue
            # 목표 거리 부근을 촘촘히 다시 읽어 실제 프레임을 고른다
            picked = 0
            for dt in np.arange(-4.0, 4.01, 0.5):
                if picked >= args.per_defect:
                    break
                t = t0 + dt
                if t < 0 or (fps and t * fps >= nf):
                    continue
                cap.set(cv2.CAP_PROP_POS_FRAMES, int(t * fps))
                ok, fr = cap.read()
                if not ok:
                    continue
                dv = read_distance(ocr_distance_from_frame, fr, tmp)
                if dv is None or abs(dv - d.pos_m) > args.tol_m:
                    continue
                fn = (f"{job.pipe}__{d.code or 'UNK'}__{d.pos_m:.1f}m"
                      f"__t{int(t*1000):07d}.jpg").replace("/", "-")
                cv2.imencode(".jpg", fr, [int(cv2.IMWRITE_JPEG_QUALITY), 92])[1].tofile(
                    str(frames_dir / fn))
                rows.append({
                    "file": fn, "pipe": job.pipe, "video": job.path.name,
                    "kind": "defect", "code": d.code, "grade": d.grade,
                    "raw_label": d.raw, "doc_pos_m": f"{d.pos_m:.1f}",
                    "ocr_pos_m": f"{dv:.1f}", "time_s": f"{t:.2f}",
                    "needs_review": int(d.needs_review),
                })
                picked += 1
                got.append(d.pos_m)
            if picked == 0:
                print(f"      - {d.pos_m:.1f}m {d.raw[:20]} : 해당 거리 프레임 못찾음")

        # 결함에서 충분히 떨어진 정상 프레임 (음성 샘플)
        far = [p for p in pts if all(abs(p[1] - g) > 3.0 for g in got)]
        for t, dv in far[:: max(1, len(far) // max(1, args.normals_per_video))][
                :args.normals_per_video]:
            cap.set(cv2.CAP_PROP_POS_FRAMES, int(t * fps))
            ok, fr = cap.read()
            if not ok:
                continue
            fn = f"{job.pipe}__NORMAL__{dv:.1f}m__t{int(t*1000):07d}.jpg".replace("/", "-")
            cv2.imencode(".jpg", fr, [int(cv2.IMWRITE_JPEG_QUALITY), 92])[1].tofile(
                str(frames_dir / fn))
            rows.append({
                "file": fn, "pipe": job.pipe, "video": job.path.name,
                "kind": "normal", "code": "", "grade": "", "raw_label": "",
                "doc_pos_m": "", "ocr_pos_m": f"{dv:.1f}", "time_s": f"{t:.2f}",
                "needs_review": 0,
            })
        cap.release()

    if tmp.exists():
        tmp.unlink()

    man = out / "manifest.csv"
    with open(man, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=[
            "file", "pipe", "video", "kind", "code", "grade", "raw_label",
            "doc_pos_m", "ocr_pos_m", "time_s", "needs_review"])
        w.writeheader()
        w.writerows(rows)

    nd = sum(1 for r in rows if r["kind"] == "defect")
    print(f"\n결함 프레임 {nd} / 정상 프레임 {len(rows)-nd} = 총 {len(rows)}장")
    print(f"관로 {len({r['pipe'] for r in rows})}개")
    print(f"-> {frames_dir}\n-> {man}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
